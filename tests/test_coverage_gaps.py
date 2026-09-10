from __future__ import annotations

import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from khz_workstation.data_service import DataWorkspaceService
from khz_workstation.models import NetworkMode
from khz_workstation.office.libreoffice import LibreOfficeEngine
from khz_workstation.office.onlyoffice import OnlyOfficeDesktopEngine
from khz_workstation.office.registry import OfficeRegistry
from khz_workstation.security.network import NetworkDenied, NetworkPolicy
from khz_workstation.security.session import SessionLockService
from khz_workstation.store import WorkspaceStore
from khz_workstation.workspace import Workspace


class StoreCoverageTests(unittest.TestCase):
    def _store(self, td: Path) -> WorkspaceStore:
        return WorkspaceStore(Path(td) / "db.sqlite", "w")

    def test_remove_item_with_descendants(self):
        with tempfile.TemporaryDirectory() as td:
            s = self._store(Path(td))
            s.upsert_item("pack/a.txt", "file")
            s.upsert_item("pack/b.txt", "file")
            s.upsert_item("other.txt", "file")
            removed = s.remove_item("pack", descendants=True)
            self.assertEqual(removed, 2)
            self.assertEqual(
                [r["relative_path"] for r in s.list_items()],
                ["other.txt"],
            )

    def test_create_data_table_valid_then_list(self):
        with tempfile.TemporaryDirectory() as td:
            s = self._store(Path(td))
            tid = s.create_data_table("T", [("A", "text"), ("B", "integer")])
            self.assertTrue(tid)
            tables = s.list_data_tables()
            self.assertEqual(len(tables), 1)
            self.assertEqual(tables[0]["name"], "T")

    def test_create_data_table_rejects_bad_name_and_empty_columns(self):
        with tempfile.TemporaryDirectory() as td:
            s = self._store(Path(td))
            with self.assertRaises(ValueError):
                s.create_data_table("1bad", [("A", "text")])
            with self.assertRaises(ValueError):
                s.create_data_table("T", [])
            with self.assertRaises(ValueError):
                s.create_data_table("T", [("A", "BOGUS")])

    def test_add_data_row_and_query_filters_sort_descending(self):
        with tempfile.TemporaryDirectory() as td:
            s = self._store(Path(td))
            tid = s.create_data_table("T", [("A", "text"), ("B", "integer")])
            s.add_data_row(tid, {"A": "x", "B": 1})
            s.add_data_row(tid, {"A": "y", "B": 2})
            _cols, rows = s.query_data(tid, sort_by="B", descending=True)
            self.assertEqual(rows[0]["B"], 2)
            _cols, filtered = s.query_data(tid, filters={"A": "x"})
            self.assertEqual(len(filtered), 1)

    def test_add_data_row_unknown_table_and_unknown_column(self):
        with tempfile.TemporaryDirectory() as td:
            s = self._store(Path(td))
            with self.assertRaises(KeyError):
                s.add_data_row("nope", {"A": 1})
            tid = s.create_data_table("T", [("A", "integer")])
            with self.assertRaises(ValueError):
                s.add_data_row(tid, {"Z": 1})

    def test_query_data_rejects_unknown_filter_and_sort_columns(self):
        with tempfile.TemporaryDirectory() as td:
            s = self._store(Path(td))
            tid = s.create_data_table("T", [("A", "integer")])
            with self.assertRaises(ValueError):
                s.query_data(tid, filters={"Z": 1})
            with self.assertRaises(ValueError):
                s.query_data(tid, sort_by="Z")


class DataServiceCoverageTests(unittest.TestCase):
    def _ws(self, td: Path) -> Workspace:
        return Workspace.create(Path(td) / "workspace", "Data")

    def test_import_xlsx_and_export_xlsx(self):
        with tempfile.TemporaryDirectory() as td:
            ws = self._ws(td)
            src = Path(td) / "in.xlsx"
            from openpyxl import Workbook
            wb = Workbook()
            sh = wb.active
            sh.title = "Data"
            sh.append(["Name", "Count"])
            sh.append(["Ops", "3"])
            sh.append(["Research", "7"])
            wb.save(src)
            table_id = DataWorkspaceService(ws).import_xlsx(src, "Counts")
            dest = Path(td) / "out.xlsx"
            DataWorkspaceService(ws).export_xlsx(table_id, dest)
            self.assertTrue(dest.is_file())
            from openpyxl import load_workbook
            reread = load_workbook(dest, read_only=True, data_only=True)
            self.assertEqual(reread.active["B1"].value, "Name")

    def test_import_xlsx_named_sheet(self):
        with tempfile.TemporaryDirectory() as td:
            ws = self._ws(td)
            src = Path(td) / "sheets.xlsx"
            from openpyxl import Workbook
            wb = Workbook()
            wb.active.title = "First"
            extra = wb.create_sheet("Second")
            extra.append(["V"])
            extra.append(["1"])
            wb.save(src)
            table_id = DataWorkspaceService(ws).import_xlsx(src, "S", sheet_name="Second")
            _cols, rows = ws.store.query_data(table_id)
            self.assertEqual(len(rows), 1)

    def test_import_csv_empty_header_row_raises(self):
        with tempfile.TemporaryDirectory() as td:
            ws = self._ws(td)
            src = Path(td) / "empty.csv"
            src.write_text("\n", encoding="utf-8")
            with self.assertRaises(ValueError):
                DataWorkspaceService(ws).import_csv(src, "E")

    def test_import_xlsx_empty_header_row_raises(self):
        with tempfile.TemporaryDirectory() as td:
            ws = self._ws(td)
            src = Path(td) / "empty.xlsx"
            from openpyxl import Workbook
            wb = Workbook()
            wb.active.append([])
            wb.save(src)
            with self.assertRaises(ValueError):
                DataWorkspaceService(ws).import_xlsx(src, "E")

    def test_import_csv_enforces_source_size_limit(self):
        with tempfile.TemporaryDirectory() as td:
            ws = self._ws(td)
            src = Path(td) / "big.csv"
            src.write_text("A\n1\n", encoding="utf-8")
            with mock.patch(
                "khz_workstation.data_service.Path.stat",
                return_value=mock.Mock(st_size=33 * 1024 * 1024),
            ):
                with self.assertRaises(ValueError):
                    DataWorkspaceService(ws).import_csv(src, "Big")

    def test_export_csv_writes_rows_and_audits(self):
        with tempfile.TemporaryDirectory() as td:
            ws = self._ws(td)
            src = Path(td) / "in.csv"
            src.write_text("A,B\n1,2\n3,4\n", encoding="utf-8")
            table_id = DataWorkspaceService(ws).import_csv(src, "T")
            dest = Path(td) / "out.csv"
            DataWorkspaceService(ws).export_csv(table_id, dest)
            self.assertIn("A,B", dest.read_text(encoding="utf-8-sig"))


class OfficeRegistryCoverageTests(unittest.TestCase):
    def test_statuses_lists_both_engines(self):
        registry = OfficeRegistry()
        infos = registry.statuses()
        names = {i.engine for i in infos}
        self.assertIn("LibreOffice", names)
        self.assertIn("ONLYOFFICE Desktop Editors", names)

    def test_convert_to_pdf_raises_when_no_pdf_engine(self):
        registry = OfficeRegistry()
        registry.engines = [OnlyOfficeDesktopEngine()]
        with self.assertRaises(FileNotFoundError):
            registry.convert_to_pdf(Path("x.docx"), Path("out"))

    def test_open_registered_or_system_uses_editable_engine(self):
        registry = OfficeRegistry()
        fake = mock.Mock()
        fake.info().can_edit = True
        fake.info().available = True
        fake.info().engine = "LibreOffice"
        fake.info().can_convert_pdf = False
        fake.open_for_edit.return_value = 4242
        registry.engines = [fake]
        self.assertEqual(registry.open_registered_or_system(Path("x.docx")), 4242)

    def test_open_registered_or_system_no_engine_non_windows_raises(self):
        registry = OfficeRegistry()
        nope = mock.Mock()
        nope.info().can_edit = False
        nope.info().available = False
        nope.info().engine = "NONE"
        nope.info().can_convert_pdf = False
        registry.engines = [nope]
        with mock.patch("khz_workstation.office.registry.os.name", "posix"):
            with self.assertRaises(FileNotFoundError):
                registry.open_registered_or_system(Path("x.docx"))


class LibreOfficeEngineCoverageTests(unittest.TestCase):
    def test_info_with_mocked_version(self):
        engine = LibreOfficeEngine(executable=Path("/bin/soffice"))
        with mock.patch("khz_workstation.office.libreoffice.subprocess.run") as run:
            run.return_value = subprocess.CompletedProcess(
                args=[], returncode=0, stdout="LibreOffice 7.6", stderr=""
            )
            info = engine.info()
        self.assertTrue(info.available)
        self.assertEqual(info.version, "LibreOffice 7.6")
        self.assertTrue(info.can_convert_pdf)

    def test_info_with_mocked_version_failure(self):
        engine = LibreOfficeEngine(executable=Path("/bin/soffice"))
        with mock.patch("khz_workstation.office.libreoffice.subprocess.run", side_effect=OSError("nope")):
            info = engine.info()
        self.assertTrue(info.available)
        self.assertIsNone(info.version)

    def test_info_when_executable_missing(self):
        engine = LibreOfficeEngine(executable=None)
        info = engine.info()
        self.assertFalse(info.available)
        self.assertFalse(info.can_convert_pdf)

    def test_open_for_edit_missing_executable_and_missing_file(self):
        engine = LibreOfficeEngine(executable=None)
        with self.assertRaises(FileNotFoundError):
            engine.open_for_edit(Path("/tmp/missing"))
        engine = LibreOfficeEngine(executable=Path("/bin/soffice"))
        with self.assertRaises(FileNotFoundError):
            engine.open_for_edit(Path("/tmp/definitely-missing-file"))

    def test_convert_to_pdf_missing_executable(self):
        engine = LibreOfficeEngine(executable=None)
        with self.assertRaises(FileNotFoundError):
            engine.convert_to_pdf(Path("x.docx"), Path("out"))

    def test_convert_to_pdf_nonzero_returncode(self):
        engine = LibreOfficeEngine(executable=Path("/bin/soffice"))
        with mock.patch("khz_workstation.office.libreoffice.subprocess.run") as run:
            run.return_value = subprocess.CompletedProcess(
                args=[], returncode=1, stdout="", stderr="boom"
            )
            with self.assertRaises(RuntimeError):
                engine.convert_to_pdf(Path("/tmp/exists.docx"), Path("/tmp/out"))

    def test_convert_to_pdf_missing_output(self):
        engine = LibreOfficeEngine(executable=Path("/bin/soffice"))
        with mock.patch("khz_workstation.office.libreoffice.subprocess.run") as run:
            run.return_value = subprocess.CompletedProcess(
                args=[], returncode=0, stdout="ok", stderr=""
            )
            with tempfile.TemporaryDirectory() as td:
                # source must exist() for the call to reach the post-check
                src = Path(td) / "x.docx"
                src.write_text("x", encoding="utf-8")
                with self.assertRaises(RuntimeError):
                    engine.convert_to_pdf(src, Path(td) / "out")


class OnlyOfficeEngineCoverageTests(unittest.TestCase):
    def test_open_for_edit_missing_executable(self):
        engine = OnlyOfficeDesktopEngine()
        engine.executable = None
        with self.assertRaises(FileNotFoundError):
            engine.open_for_edit(Path("x"))

    def test_convert_to_pdf_not_implemented(self):
        engine = OnlyOfficeDesktopEngine()
        with self.assertRaises(NotImplementedError):
            engine.convert_to_pdf(Path("x"), Path("out"))

    def test_init_detects_executable_via_which(self):
        with mock.patch("khz_workstation.office.onlyoffice.shutil.which") as which:
            which.side_effect = lambda name: "/bin/sh" if name == "DesktopEditors" else None
            engine = OnlyOfficeDesktopEngine()
        self.assertIsNotNone(engine.executable)
        info = engine.info()
        self.assertTrue(info.available)

    def test_open_for_edit_invokes_subprocess_when_executable_present(self):
        engine = OnlyOfficeDesktopEngine()
        engine.executable = Path("/usr/bin/DesktopEditors")
        with mock.patch("khz_workstation.office.onlyoffice.subprocess.Popen") as popen:
            popen.return_value = mock.Mock(pid=99)
            self.assertEqual(engine.open_for_edit(Path("doc.docx")), 99)


class NetworkPolicyCoverageTests(unittest.TestCase):
    def test_unrestricted_allows_any_host(self):
        self.assertTrue(NetworkPolicy(mode=NetworkMode.UNRESTRICTED).authorize_host("evil.example.com"))

    def test_allowlist_matches_host(self):
        policy = NetworkPolicy(mode=NetworkMode.ALLOWLIST, allowlist=("allowed.example.com",))
        self.assertTrue(policy.authorize_host("ALLOWED.example.com"))
        self.assertFalse(policy.authorize_host("denied.example.com"))

    def test_loopback_only_accepts_loopback_ip_literal(self):
        policy = NetworkPolicy(mode=NetworkMode.LOOPBACK_ONLY)
        self.assertTrue(policy.authorize_host("127.0.0.1"))
        self.assertFalse(policy.authorize_host("8.8.8.8"))

    def test_authorize_url_raises_for_denied(self):
        policy = NetworkPolicy(mode=NetworkMode.DENY)
        with self.assertRaises(NetworkDenied):
            policy.authorize_url("https://evil.example.com/path")

    def test_authorize_url_raises_for_missing_host(self):
        policy = NetworkPolicy(mode=NetworkMode.DENY)
        with self.assertRaises(NetworkDenied):
            policy.authorize_url("not-a-url")

    def test_loopback_only_dns_resolved_to_loopback(self):
        policy = NetworkPolicy(mode=NetworkMode.LOOPBACK_ONLY)
        with mock.patch("khz_workstation.security.network.socket.getaddrinfo") as gai:
            gai.return_value = [(None, None, None, None, ("127.0.0.1", 0))]
            self.assertTrue(policy.authorize_host("localhost"))

    def test_loopback_only_dns_failure_denies(self):
        policy = NetworkPolicy(mode=NetworkMode.LOOPBACK_ONLY)
        with mock.patch("khz_workstation.security.network.socket.getaddrinfo", side_effect=OSError("nope")):
            self.assertFalse(policy.authorize_host("host.example.com"))


class SessionLockCoverageTests(unittest.TestCase):
    def test_supported_with_lock_impl(self):
        svc = SessionLockService(lock_impl=lambda: True)
        self.assertTrue(svc.supported)

    def test_lock_now_invokes_impl(self):
        svc = SessionLockService(lock_impl=lambda: True)
        self.assertTrue(svc.lock_now())

    def test_lock_now_returns_false_without_impl_on_posix(self):
        svc = SessionLockService()
        with mock.patch("khz_workstation.security.session.os.name", "posix"):
            self.assertFalse(svc.lock_now())


if __name__ == "__main__":
    unittest.main()
